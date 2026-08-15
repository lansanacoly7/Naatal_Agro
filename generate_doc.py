import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

def generate_doc():
    doc = Document()

    # --- PAGE DE GARDE ---
    inst_para = doc.add_paragraph()
    inst_run = inst_para.add_run("Institut Supérieur d'Enseignement Professionnel (ISEP)")
    inst_run.font.size = Pt(16)
    inst_run.bold = True
    inst_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("\n\n")

    proj_para = doc.add_paragraph()
    proj_run = proj_para.add_run("Rapport de Projet de Fin de Formation")
    proj_run.font.size = Pt(20)
    proj_run.bold = True
    proj_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("\n")

    title_para = doc.add_paragraph()
    title_run = title_para.add_run("NAATAL AGRO : Copilote Agricole Intelligent pour le Sénégal")
    title_run.font.size = Pt(24)
    title_run.bold = True
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("\n\n\n\n\n\n")

    author_para = doc.add_paragraph()
    author_run = author_para.add_run("Réalisé par :\nLassana Coly")
    author_run.font.size = Pt(14)
    author_run.bold = True
    author_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_page_break()

    # --- TABLE DES MATIERES ---
    toc_title = doc.add_heading("Table des matières (Adaptée à Naatal Agro)", level=1)
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # Introduction
    doc.add_paragraph("Introduction générale")
    doc.add_paragraph("1. Contexte et justification du sujet : L'agriculture au Sénégal face aux défis climatiques et technologiques")
    doc.add_paragraph("2. Problématique : Comment centraliser les informations agricoles pour optimiser la prise de décision ?")
    doc.add_paragraph("3. Questions de recherche, générale et spécifiques")
    doc.add_paragraph("4. Objectif général et objectifs subsidiaires : Création de la plateforme SaaS Naatal Agro")
    doc.add_paragraph("5. Hypothèses de recherche : La donnée centralisée et l'IA améliorent le rendement")
    doc.add_paragraph("6. Méthodologie retenue : Méthode Agile et Clean Architecture")
    doc.add_paragraph("7. Annonce du plan")

    doc.add_paragraph()

    # Chapitre 1
    doc.add_heading("CHAPITRE I : CADRE THÉORIQUE ET CONCEPTUEL", level=2)
    doc.add_paragraph("Section 1 : Revue conceptuelle")
    doc.add_paragraph("    Sous-section 1 : L'agriculture de précision et les solutions SaaS agricoles")
    doc.add_paragraph("    Sous-section 2 : Le rôle de l'Intelligence Artificielle et des données météorologiques")
    doc.add_paragraph("Section 2 : Revue de littérature et état de l'art (Solutions existantes au Sénégal)")

    doc.add_paragraph()

    # Chapitre 2
    doc.add_heading("CHAPITRE II : CONTEXTE DE L'ÉTUDE", level=2)
    doc.add_paragraph("Section 1 : Contexte général")
    doc.add_paragraph("    Sous-section 1 : L'importance du secteur agricole dans l'économie sénégalaise")
    doc.add_paragraph("    Sous-section 2 : Les défis majeurs : accès à l'information et aux marchés")
    doc.add_paragraph("Section 2 : Contexte spécifique (Le projet Naatal Agro)")
    doc.add_paragraph("    Sous-section 1 : Genèse et objectifs de la plateforme Naatal Agro")
    doc.add_paragraph("    Sous-section 2 : Public cible (Agriculteurs, Coopératives) et leurs attentes")

    doc.add_paragraph()

    # Chapitre 3
    doc.add_heading("CHAPITRE III : CADRE MÉTHODOLOGIQUE ET CONCEPTION", level=2)
    doc.add_paragraph("Section 1 : Spécification et modélisation du système")
    doc.add_paragraph("    Sous-section 1 : Spécification des besoins fonctionnels (Auth, Dashboard, Cultures, IA) et non-fonctionnels (Sécurité JWT, Scalabilité)")
    doc.add_paragraph("    Sous-section 2 : Étude et choix des méthodes de modélisation (Clean Architecture, Modélisation relationnelle)")
    doc.add_paragraph("    Sous-section 3 : Modélisation (Schéma de la base de données PostgreSQL, Diagrammes de flux)")
    doc.add_paragraph("Section 2 : Implémentation et évaluation des modèles")
    doc.add_paragraph("    Sous-section 1 : Choix technologiques et justification (Flutter pour le mobile, Django REST Framework pour le backend)")
    doc.add_paragraph("    Sous-section 2 : Description de l'expérimentation (Développement de l'onboarding, intégration Auth JWT, réalisation du Dashboard dynamique)")

    # Save
    save_path = os.path.abspath('Table_Des_Matieres_Naatal_Agro.docx')
    doc.save(save_path)
    print(f"Document Word généré avec succès : {save_path}")

if __name__ == '__main__':
    generate_doc()
