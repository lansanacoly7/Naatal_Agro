"""
Assemble les fichiers Markdown du mémoire (01-*.md, 02-*.md, 03-*.md) en un document Word.

Usage :  python construire_docx.py
Résultat : ../Memoire_Naatal_Agro_brouillon.docx

Les titres utilisent les styles « Titre 1/2/3 » : dans Word, clic droit sur la table des matières,
« Mettre à jour les champs », pour obtenir les numéros de page. Les passages [À COMPLÉTER ...] et
[À VÉRIFIER ...] sont surlignés en jaune.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
OUTPUT = HERE.parent / 'Memoire_Naatal_Agro_brouillon.docx'
PARTS = sorted(HERE.glob('0[1-9]-*.md'))

TODO = re.compile(r'(\[(?:À COMPLÉTER|À VÉRIFIER|Insérer|Pistes)[^\]]*\])')
INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)')


def add_runs(paragraph, text):
    """Ajoute du texte avec gras, italique, code et surlignage des passages à compléter."""
    for chunk in TODO.split(text):
        if not chunk:
            continue
        highlight = bool(TODO.fullmatch(chunk))
        for piece in INLINE.split(chunk):
            if not piece:
                continue
            if piece.startswith('**') and piece.endswith('**'):
                run = paragraph.add_run(piece[2:-2])
                run.bold = True
            elif piece.startswith('`') and piece.endswith('`'):
                run = paragraph.add_run(piece[1:-1])
                run.font.name = 'Consolas'
            elif piece.startswith('*') and piece.endswith('*') and len(piece) > 2:
                run = paragraph.add_run(piece[1:-1])
                run.italic = True
            else:
                run = paragraph.add_run(piece)
            if highlight:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def set_cell_shading(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), color)
    properties.append(shading)


def add_table(document, rows):
    header, body = rows[0], rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ''
        add_runs(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
        set_cell_shading(cell, 'DDEBDD')
    for row in body:
        cells = table.add_row().cells
        for index in range(len(header)):
            cells[index].text = ''
            add_runs(cells[index].paragraphs[0], row[index] if index < len(row) else '')
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(10)
    document.add_paragraph()


def add_toc(document):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t')
    placeholder.text = 'Clic droit ici, puis « Mettre à jour les champs » pour afficher la table des matières.'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for element in (begin, instruction, separate, placeholder, end):
        run._r.append(element)


def cover_page(document):
    def centered(text, size, bold=False, space_after=12):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(space_after)
        add_runs(paragraph, text)
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.bold = bold

    centered("Institut Supérieur d'Enseignement Professionnel (ISEP)", 16, True, 36)
    centered('Rapport de Projet de Fin de Formation', 14, False, 48)
    centered('NAATAL AGRO', 28, True, 8)
    centered("Application mobile d'aide à la décision pour les producteurs sénégalais, avec un assistant "
             'conversationnel aux réponses sourcées', 15, False, 60)
    centered('Réalisé par : Lassana Coly', 13, False, 6)
    centered('[À COMPLÉTER : filière, niveau, année académique]', 12, False, 6)
    centered('[À COMPLÉTER : encadrant pédagogique, encadrant professionnel]', 12, False, 6)
    centered('[À COMPLÉTER : date de soutenance et mois de dépôt]', 12, False, 6)
    document.add_page_break()
    for title in ('Remerciements', 'Résumé', 'Abstract'):
        document.add_paragraph(title, style='Heading 1')
        add_runs(document.add_paragraph(), '[À COMPLÉTER : à rédiger en dernier, une fois le contenu validé.]')
    document.add_page_break()
    document.add_paragraph('Table des matières', style='Heading 1')
    add_toc(document)
    document.add_page_break()


def convert(document, markdown):
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(document, table_rows)
            table_rows = []

    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith('|'):
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                continue
            table_rows.append(cells)
            continue
        flush_table()
        if not line.strip() or line.strip() == '---':
            continue
        heading = re.match(r'^(#{1,3})\s+(.*)$', line)
        if heading:
            level, title = len(heading.group(1)), heading.group(2)
            if level == 1 and title.startswith(('CHAPITRE', 'Conclusion', 'Webographie')):
                document.add_page_break()
            document.add_paragraph(title, style=f'Heading {level}')
            continue
        quote = re.match(r'^>\s*(.*)$', line)
        if quote:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1)
            add_runs(paragraph, quote.group(1))
            for run in paragraph.runs:
                run.italic = True
            continue
        bullet = re.match(r'^[-*]\s+(.*)$', line)
        if bullet:
            add_runs(document.add_paragraph(style='List Bullet'), bullet.group(1))
            continue
        numbered = re.match(r'^\d+\.\s+(.*)$', line)
        if numbered:
            add_runs(document.add_paragraph(style='List Number'), numbered.group(1))
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(paragraph, line)
    flush_table()


def main():
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.5)
    normal = document.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12)):
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True

    cover_page(document)
    for part in PARTS:
        convert(document, part.read_text(encoding='utf-8'))
    document.save(OUTPUT)
    words = sum(len(part.read_text(encoding='utf-8').split()) for part in PARTS)
    print(f'{OUTPUT.name} généré ({words} mots, {len(PARTS)} fichiers).')


if __name__ == '__main__':
    sys.exit(main())
